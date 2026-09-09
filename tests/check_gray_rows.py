"""Regression checks for gray exclusions, conditional fills, ranking and export.

Run: python tests/check_gray_rows.py
Requires selenium, Chrome and openpyxl. Uses temporary synthetic workbooks only.
Optional --rating PATH independently checks conditional colors in a local source
book without printing employee data or copying the workbook into the repository.
"""
import argparse
import json
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Color, PatternFill
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait

SITE = Path(__file__).resolve().parents[1]
HEADERS = ['ПІБ КС', 'набрано балів', 'перегляд', 'години (графік)', 'поза рейтингом',
           'години в рейтинг', 'бали на годину', 'ефективність', 'рейтингу', 'РЗ',
           'РЗ в рейтингу', 'КК', 'КК в рейтингу', 'Фінал', 'Порядок у джерелі']


def create_inputs(folder):
    rating = Workbook()
    ws = rating.active
    ws.title = 'Рейтинг фінал (Тест)'
    ws['K1'] = '=MAX(G3:G13)'
    ws.append(HEADERS)
    people = [
        ('Active Leader', 10, 10, 1), ('Active Second', 5, 10, 2),
        ('Gray Conditional', 1000, 10, 57), ('Gray ZeroHours', 100, 0, 4),
        ('Gray Theme', 200, 10, 5), ('Gray Indexed', 300, 10, 6),
        ('Active Override', 3, 10, 777), ('Active Stripe', 2, 10, 8),
        ('Active SingleCell', 1, 10, 9), ('Active FewHours', .05, 1, 10),
        ('Active Tie', 10, 10, 11),
    ]
    gray = PatternFill('solid', fgColor='FFB7B7B7')
    for row_number, (name, points, hours, source_rank) in enumerate(people, start=3):
        ws.append([name, points, 0, hours, 0, f'=D{row_number}-E{row_number}',
                   f'=(B{row_number}+C{row_number})/F{row_number}',
                   f'=G{row_number}/$K$1*100', f'=H{row_number}*0.4', 100, 20, 90, 36, 999, source_rank])
        fill = None
        if name in ('Gray ZeroHours', 'Active Override'):
            fill = gray
        elif name == 'Gray Theme':
            fill = PatternFill('solid', fgColor=Color(theme=1, tint=.7))
        elif name == 'Gray Indexed':
            fill = PatternFill('solid', fgColor=Color(indexed=23))
        elif name == 'Active Stripe':
            fill = PatternFill('solid', fgColor='FFF9F9F9')
        if fill:
            for cell in ws[row_number][:14]:
                cell.fill = fill
        if name == 'Active SingleCell':
            ws.cell(row_number, 2).fill = gray
    ws.conditional_formatting.add('A3:N13', FormulaRule(formula=['$O:$O=777'],
                                 fill=PatternFill('solid', fgColor='FF70AD47')))
    ws.conditional_formatting.add('A3:N13', FormulaRule(formula=['$O:$O=57'], fill=gray))
    rating.save(folder / 'rating.xlsx')
    rating.close()

    quality = Workbook()
    quality.active.append(['ПІБ', 'Середній бал'])
    for name, *_ in people:
        quality.active.append([name, 90])
    quality.save(folder / 'statistics.xlsx')
    quality.close()


CORE_CHECKS = r'''
const c = RatingCore;
let checks = 0;
function ok(test, label) { if (!test) throw new Error(label); checks++; }
const rows = state.originalResults, ranked = state.rankedResults;
const grays = rows.filter(row => row.excludedFromRating);
ok(rows.length === 11 && grays.length === 4, 'four actual gray rows');
ok(grays.every(row => row.fio.startsWith('Gray ')), 'conditional, solid, theme, indexed');
ok(rows.filter(row => row.fio.startsWith('Active ')).every(row => !row.excludedFromRating),
   'colored override, near-white, isolated numeric fill and few hours stay eligible');
ok(state.maximum === 1, 'gray high scorer and zero hours excluded from MAX');
ok(ranked.slice(0, 7).every(row => !row.excludedFromRating), 'participants first');
ok(ranked.slice(7).every(row => row.place === null && row.zone.css === 'zone-unrated'), 'grays last without places');
ok(grays.find(row => row.fio === 'Gray Conditional').finalScore > ranked[0].finalScore, 'high gray score retained for reference');
ok(ranked[0].place === 1 && ranked[1].place === 1 && ranked[2].place === 3, 'ties among participants');
ok(ranked[0].zone.css === 'zone-green' && ranked[1].zone.css === 'zone-green'
   && ranked[2].zone.css === 'zone-yellow' && ranked[6].zone.css === 'zone-red', 'zones use seven eligible people only');
ok(document.querySelectorAll('#ratingBody tr.zone-unrated').length === 4, 'UI grays');
ok(document.getElementById('participationSummary').textContent.includes('У рейтингу: 7.'), 'participant count');
const bonusRows = rows.map(row => ({...row, issues: []}));
c.applyReviews(bonusRows, {entries: [{fio:'Gray Conditional',date:{period:'2026-09'},decision:'так',points:999}],sheet:'Перегляди'}, '2026-09');
c.recalculate(bonusRows, {first:3,last:13});
ok(rankResults(bonusRows).slice(7).every(row => row.place === null), 'review bonus cannot requalify gray employee');

function probe(formula, marker = 57) {
  const wb = new ExcelJS.Workbook(), ws = wb.addWorksheet('Probe');
  ws.addRow(c.columns.map(x => x.title).concat('Source rank'));
  ws.addRow(['Person Example',10,0,10,0,10,1,100,40,0,0,90,36,76,marker]);
  ws.addConditionalFormatting({ref:'A2:N2',rules:[{type:'expression',priority:2,formulae:[formula],
    style:{fill:{type:'pattern',pattern:'solid',fgColor:{argb:'FFB7B7B7'}}}}]});
  return {ws, read:() => c.readRows(ws,c.ratingLayout(ws))};
}
ok(probe('$O2=57').read()[0].excludedFromRating, 'relative row expression');
ok(!probe('$O:$O=58').read()[0].excludedFromRating, 'false gray rule');
ok(probe('TRUE()').read()[0].excludedFromRating, 'unconditional gray rule');
ok(!probe('FALSE').read()[0].excludedFromRating, 'false constant');
ok(probe('$O:$O=0', {formula:'1-1',result:0}).read()[0].excludedFromRating, 'cached zero evaluated');
for (const test of [probe('MOD(ROW(),2)=0'),probe('$O:$O=57',{formula:'1+56'})]) {
  let stopped=false;
  try { test.read(); } catch(e) { stopped=e.message.includes('сіре позначення'); }
  ok(stopped, 'unknown rule or missing cache must not silently qualify an employee');
}
const priority = probe('TRUE()');
priority.ws.conditionalFormattings[0].rules.unshift({type:'expression',priority:1,stopIfTrue:true,formulae:['TRUE()'],style:{font:{bold:true}}});
ok(!priority.read()[0].excludedFromRating, 'stopIfTrue respected');
const relative = probe('$O2=57');
relative.ws.addRow(['Other Example',10,0,10,0,10,1,100,40,0,0,90,36,76,1]);
relative.ws.conditionalFormattings[0].ref='A2:N3';
ok(relative.read()[0].excludedFromRating && !relative.read()[1].excludedFromRating, 'relative reference shifts by row');
const allGray = rows.map(row => ({...row,excludedFromRating:true,issues:[]}));
ok(c.recalculate(allGray,{first:3,last:13}) === null, 'all gray book has no artificial base');
ok(rankResults(allGray).every(row => row.place === null && row.finalScore === null), 'all gray no places or divided-by-zero');
return checks;
'''


def check_source(driver, wait, path):
    """Independent oracle for this source's $O:$O comparisons and solid fills."""
    import re
    import xml.etree.ElementTree as ET
    from zipfile import ZipFile
    from openpyxl.utils.cell import range_boundaries
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    expected = []
    with ZipFile(path) as archive:
        styles = ET.fromstring(archive.read('xl/styles.xml'))
        fills, xfs, dxfs = [list(styles.find(f's:{tag}', ns)) for tag in ('fills', 'cellXfs', 'dxfs')]
        relations = {r.attrib['Id']: r.attrib['Target'] for r in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        sheets = ET.fromstring(archive.read('xl/workbook.xml')).find('s:sheets', ns)
        for sheet in sheets:
            name = sheet.attrib['name']
            if not name.lower().startswith('рейтинг фінал'):
                continue
            target = relations[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
            target = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            doc = ET.fromstring(archive.read(target))
            excluded = []
            for row in doc.findall('s:sheetData/s:row', ns):
                row_number = int(row.attrib['r'])
                cells = {c.attrib['r']: c for c in row}
                cell = cells.get(f'A{row_number}')
                if row_number < 3 or cell is None or cell.find('s:v', ns) is None:
                    continue
                fill = fills[int(xfs[int(cell.attrib.get('s', 0))].attrib.get('fillId', 0))]
                rules = []
                for formatting in doc.findall('s:conditionalFormatting', ns):
                    ranges = [range_boundaries(ref) for ref in formatting.attrib['sqref'].split()]
                    if not any(left <= 1 <= right and top <= row_number <= bottom for left, top, right, bottom in ranges):
                        continue
                    rules.extend(formatting)
                for rule in sorted(rules, key=lambda r: int(r.attrib['priority'])):
                    dxf = dxfs[int(rule.attrib['dxfId'])].find('s:fill', ns)
                    if dxf is None:
                        continue
                    match = re.fullmatch(r'\$O:\$O=(\d+)', rule.find('s:formula', ns).text)
                    assert match, 'Source changed: extend independent formatting oracle'
                    rank_cell = cells.get(f'O{row_number}')
                    rank = rank_cell.find('s:v', ns) if rank_cell is not None else None
                    if rank is not None and float(rank.text) == int(match[1]):
                        fill = dxf
                        break
                pattern = fill.find('s:patternFill', ns)
                color = pattern.find('s:fgColor', ns) if pattern is not None else None
                if pattern is not None and pattern.attrib.get('patternType') == 'solid' and color is not None:
                    if color.attrib.get('rgb', '')[-6:] in ('B7B7B7', 'CCCCCC', 'D9D9D9'):
                        excluded.append(row_number)
            expected.append((name, excluded))
    driver.get((SITE / 'index.html').as_uri())
    driver.find_element(By.ID, 'ratingFile').send_keys(str(path.resolve()))
    wait.until(lambda d: not d.execute_script('return state.ratingLoading'))
    assert driver.execute_script('return Boolean(state.ratingWorkbook)'), 'Source load failed'
    counts = []
    for name, excluded in expected:
        actual = driver.execute_script('''
            const ws = state.ratingWorkbook.getWorksheet(arguments[0]);
            return RatingCore.readRows(ws,RatingCore.ratingLayout(ws)).filter(r=>r.excludedFromRating).map(r=>r.sourceRow);
        ''', name)
        assert actual == excluded, f'Source exclusion mismatch; sheet index {len(counts)}'
        counts.append(len(actual))
    return counts


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--rating', type=Path)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix='rating_gray_') as tmp:
        folder = Path(tmp)
        create_inputs(folder)
        options = webdriver.ChromeOptions()
        options.add_argument('--headless=new')
        options.add_argument('--disable-background-networking')
        options.add_experimental_option('prefs', {'download.default_directory': tmp, 'download.prompt_for_download': False})
        options.set_capability('goog:loggingPrefs', {'browser': 'ALL'})
        with webdriver.Chrome(options=options) as driver:
            wait = WebDriverWait(driver, 90)
            driver.get((SITE / 'index.html').as_uri())
            driver.execute_cdp_cmd('Network.emulateNetworkConditions', {
                'offline': True, 'latency': 0, 'downloadThroughput': 0, 'uploadThroughput': 0,
            })
            driver.find_element(By.ID, 'qualityFile').send_keys(str(folder / 'statistics.xlsx'))
            driver.find_element(By.ID, 'ratingFile').send_keys(str(folder / 'rating.xlsx'))
            wait.until(lambda d: d.find_element(By.ID, 'calculateButton').is_enabled())
            driver.find_element(By.ID, 'calculateButton').click()
            wait.until(lambda d: not d.execute_script('return state.busy'))
            assert driver.find_element(By.ID, 'results').is_displayed(), driver.find_element(By.ID, 'progress').text
            checks = driver.execute_script(CORE_CHECKS)
            driver.find_element(By.ID, 'downloadExcel').click()
            exported = folder / 'current_final_rating.xlsx'
            wait.until(lambda d: exported.exists() and not list(folder.glob('*.crdownload')))
            wb = load_workbook(exported, data_only=True)
            try:
                ws = wb.worksheets[0]
                headers = [cell.value for cell in ws[1]]
                place = headers.index('Місце')
                status = headers.index('Статус КЯ / розрахунку')
                rows = list(ws.iter_rows(min_row=2))
                assert len(rows) == 11
                assert all(row[place].value is not None for row in rows[:7])
                assert all(row[place].value is None and row[0].fill.fgColor.rgb == 'FFD9D9D9' for row in rows[7:])
                assert all('Не рейтингується' in row[status].value for row in rows[7:])
                assert rows[7][13].value > rows[0][13].value
                assert any(row[0].value == 'Сіра — поза рейтингом' for row in wb['Легенда'])
            finally:
                wb.close()
            assert not [entry for entry in driver.get_log('browser') if entry['level'] == 'SEVERE']
            result = {'synthetic_checks': checks, 'eligible': 7, 'excluded': 4, 'excel_order_and_colors': True, 'offline': True}
            if args.rating:
                result['real_book_excluded_per_sheet'] = check_source(driver, wait, args.rating)
            print(json.dumps(result))


if __name__ == '__main__':
    main()
